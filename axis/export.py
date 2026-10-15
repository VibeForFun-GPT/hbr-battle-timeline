"""将行动轴导出为 PNG、Excel、CSV、HTML 与带 Evidence 的 ZIP。"""
from __future__ import annotations

import csv
from html import escape
import json
import math
from pathlib import Path
import zipfile

from PIL import Image, ImageDraw, ImageFont
from PIL.PngImagePlugin import PngInfo
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .model import compact_rows, stamp, turn_coverage

COLORS = ["#faeee5", "#eeedfa", "#e8f5ef"]
HEADERS = ["回合", "位置 1 · 角色", "位置 1 · Skill", "位置 2 · 角色", "位置 2 · Skill", "位置 3 · 角色", "位置 3 · Skill", "行动后 OD", "确认时间", "备注"]


def display_rows(document, compact=False):
    return compact_rows(document["rows"]) if compact else document["rows"]


def gap_label(gap):
    missing=gap.get("missing_turns",[])
    if missing:
        return str(missing[0]) if len(missing)==1 else f"{missing[0]}–{missing[-1]}"
    return f"{gap.get('after_turn','?')} → {gap.get('before_turn','?')} 之间"


def gap_time_range(gap):
    return " – ".join(stamp(gap[key]) if gap.get(key) is not None else "时间未记录" for key in ("start_time","end_time"))


def gap_text(gap):
    return f"缺失回合 {gap_label(gap)}（未识别） · {gap_time_range(gap)}"


def gap_notes(row,coverage):
    notes=[]
    for gap in (coverage or {}).get("gaps",[]):
        if row.get("id") and gap.get("before_row_id"):
            adjacent=row["id"]==gap["before_row_id"]
        else:
            adjacent=row.get("turn")==gap.get("before_turn")
        if adjacent:notes.append(gap_text(gap)+"；该区间的角色与 Skill 未补填")
    return notes


def fields(row, full=False, od="after",coverage=None):
    values=[row.get("label", str(row.get("turn","")))]
    for action in row["actions"]:
        values.extend([action.get("character_full" if full else "character", ""),action.get("skill_full" if full else "skill", "")])
    value=row.get("od_before" if od=="before" else "od_after")
    extras="；".join(a["character"]+"「"+a["skill"]+"」" for a in row.get("extra_actions",[]))
    notes=[]
    if row.get("needs_review") and not row.get("reviewed"):
        notes.append("待核对")
    if extras:
        notes.append("追加换人："+extras)
    if row.get("bonus_text"):
        notes.append(row["bonus_text"])
    if row.get("note"):
        notes.append(row["note"])
    notes.extend(gap_notes(row,coverage))
    return values + [value if value is not None else "",stamp(row.get("confirm_time",row["time"])),"；".join(notes)]


def excel_text(value):
    if isinstance(value,str) and value[:1] in ("=","+","-","@"):
        return "'"+value
    return value


def export_csv(document,path,*,compact=False,full=False,od="after"):
    coverage=turn_coverage(document["rows"])
    headers=HEADERS.copy()
    if od=="before":headers[7]="行动前 OD"
    with Path(path).open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f)
        writer.writerow(headers)
        for row in display_rows(document,compact):
            writer.writerow([excel_text(v) for v in fields(row,full,od,coverage)])


def export_excel(document,path,*,compact=False,od="after"):
    coverage=turn_coverage(document["rows"])
    gaps=coverage["gaps"]
    wb=Workbook()
    wb.remove(wb.active)
    headers=HEADERS.copy()
    headers[7]="行动前 OD" if od=="before" else "行动后 OD"
    for name,full in [("行动轴",False),("完整 Skill 名称",True)]:
        ws=wb.create_sheet(name)
        ws.append([document["title"]+(" · 存在缺失回合，请查看「缺失回合」" if gaps else "")])
        ws.merge_cells("A1:J1")
        ws.append(headers)
        for row in display_rows(document,compact):
            ws.append(fields(row,full,od,coverage))
            r=ws.max_row
            for c in range(1,11):
                cell=ws.cell(r,c)
                if isinstance(cell.value,str):cell.data_type="s"
                cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
                if 2<=c<=7:cell.fill=PatternFill("solid",fgColor=COLORS[(c-2)//2].lstrip("#"))
            if row.get("kind")!="normal":ws.cell(r,1).font=Font(color="F05E6A",bold=True)
            ws.cell(r,8).font=Font(color="478EC9",bold=True)
            if row.get("needs_review") and not row.get("reviewed"):ws.cell(r,10).fill=PatternFill("solid",fgColor="FFF0D0")
            ws.row_dimensions[r].height=34
        ws.freeze_panes="B3"
        ws.auto_filter.ref=f"A2:J{ws.max_row}"
        for c,width in enumerate([15,14,24,14,24,14,24,14,16,45],1):ws.column_dimensions[get_column_letter(c)].width=width
        ws.row_dimensions[1].height=30
        for cell in ws[1]+ws[2]:cell.font=Font(bold=True,color="24344D",size=12);cell.alignment=Alignment(horizontal="center",vertical="center")
        if gaps:ws.cell(1,1).fill=PatternFill("solid",fgColor="FFF0D0")
    ws=wb.create_sheet("识别依据")
    ws.append(["序号","Frame 时间","确认时间","视频 TURN","BONUS","前排 Avatar ID","完整位置 1–6","Confidence","已核对","Screenshot","识别提示"])
    for row in document["rows"]:
        ws.append([row["index"],stamp(row["time"]),stamp(row["confirm_time"]),row["turn"],row.get("bonus_text",""),",".join(str(a["character_id"]) for a in row["actions"]),",".join(map(str,row.get("positions",[]))),row.get("confidence"),"是" if row.get("reviewed") else "否",row.get("screenshot",""),"；".join(row.get("notes",[]))])
    ws.freeze_panes="A2"
    for c in range(1,12):ws.column_dimensions[get_column_letter(c)].width=20 if c<11 else 65
    if gaps:
        missing=wb.create_sheet("缺失回合")
        missing.append(["缺失回合","前一已记录 TURN","后一已记录 TURN","时间范围","状态","说明"])
        for gap in gaps:
            missing.append([gap_label(gap),gap.get("after_turn"),gap.get("before_turn"),
                            gap_time_range(gap),"未识别",
                            gap.get("message") or "该区间没有记录到可见选招；角色与 Skill 未补填，请回看视频核对。"])
            for cell in missing[missing.max_row]:
                if isinstance(cell.value,str):cell.data_type="s"
                cell.alignment=Alignment(vertical="center",wrap_text=True)
            missing.row_dimensions[missing.max_row].height=42
        for cell in missing[1]:cell.font=Font(bold=True,color="9A5B16");cell.fill=PatternFill("solid",fgColor="FFF0D0")
        for c,column_width in enumerate([16,20,20,28,14,85],1):missing.column_dimensions[get_column_letter(c)].width=column_width
        missing.freeze_panes="A2"
    about=wb.create_sheet("说明")
    for line in ["视频："+document["video"]["name"],"开始时间："+stamp(document["settings"]["start"]),"OD 数据来源：视频画面 OCR。",*document.get("warnings",[])]:about.append([line])
    for gap in gaps:about.append([gap_text(gap)+"；角色与 Skill 未补填。"])
    about.column_dimensions["A"].width=105
    wb.save(path)


def font_path():
    options=["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]
    for path in options:
        if Path(path).exists():return path
    raise ValueError("PNG 导出需要中文 Font。请在系统安装 Microsoft YaHei 或 Noto Sans CJK。")


def export_png(document,path,*,compact=False,od="after"):
    rows=display_rows(document,compact)
    coverage=turn_coverage(document["rows"])
    gaps=coverage["gaps"]
    gap_height=40+34*len(gaps) if gaps else 0
    widths=[128,118,170,118,170,118,170,127,115]
    width=sum(widths)
    header_h,row_h,footer_h=125+gap_height,53,92
    im=Image.new("RGB",(width,header_h+len(rows)*row_h+footer_h),"white")
    draw=ImageDraw.Draw(im)
    fonts=[ImageFont.truetype(font_path(),size) for size in (17,21,24)]

    def fitted(text,x,y,w,h,color="#24334b",size=21):
        text=str(text)
        font=ImageFont.truetype(font_path(),size)
        while draw.textbbox((0,0),text,font=font)[2]>w-12 and size>12:
            size-=1;font=ImageFont.truetype(font_path(),size)
        draw.text((x+w/2,y+h/2),text,fill=color,font=font,anchor="mm")

    draw.rectangle((0,0,width,75),fill="#f4f7fb")
    fitted(document["title"],0,0,width,40,size=24)
    team=document["team"]
    lineup="前："+" / ".join(p["alias"] for p in team[:3])+"  |  后："+" / ".join(p["alias"] for p in team[3:])
    fitted(lineup,0,40,width,29,size=17)
    if gaps:
        draw.rectangle((0,75,width,75+gap_height),fill="#fff0d3",outline="#e8c98f")
        fitted(f"检测到 {coverage['missing_turn_count']} 个缺失回合，以下区间未识别；角色与 Skill 未补填。",0,75,width,40,color="#945b20",size=18)
        for index,gap in enumerate(gaps):fitted(gap_text(gap),0,115+index*34,width,34,color="#945b20",size=17)
    x=0
    table_y=75+gap_height
    labels=["回合","角色","Skill","角色","Skill","角色","Skill","行动前 OD" if od=="before" else "行动后 OD","确认时间"]
    for c,w in enumerate(widths):
        draw.rectangle((x,table_y,x+w,table_y+50),fill=COLORS[(c-1)//2] if 1<=c<=6 else "#f3f5f8",outline="#d4dce6")
        fitted(labels[c],x,table_y,w,50,size=17)
        x+=w
    for r,row in enumerate(rows):
        values=fields(row,od=od)[:9]
        y=header_h+r*row_h
        x=0
        for c,(value,w) in enumerate(zip(values,widths)):
            fill=COLORS[(c-1)//2] if 1<=c<=6 else "#f8fafc"
            if c==0 and row.get("needs_review") and not row.get("reviewed"):fill="#fff0d3"
            draw.rectangle((x,y,x+w,y+row_h),fill=fill,outline="#d4dce6")
            color="#ed5b6b" if c==0 and row.get("kind")!="normal" else "#478dc9" if c==7 else "#29374c"
            fitted(value,x,y,w,row_h,color,size=22 if c!=8 else 16)
            x+=w
    y=header_h+len(rows)*row_h
    note="黄色回合格 = 待核对；OD 为视频中的可见读数；完整版 Skill 与 Screenshot 见 Excel / JSON。"
    fitted(note,0,y,width,36,size=16)
    extras=[f"{r.get('label')}："+"、".join(a["character"]+" "+a["skill"] for a in r.get("extra_actions",[])) for r in rows if r.get("extra_actions")]
    if extras:
        fitted("追加换人："+"；".join(extras),0,y+35,width,30,size=14)
    else:
        fitted("每行对应一次实际行动确认"+("；追加确认已合并" if compact else "；保留所有追加确认"),0,y+35,width,30,size=15)
    metadata=PngInfo()
    if gaps:metadata.add_text("Coverage","\n".join(gap_text(gap) for gap in gaps)+"\n角色与 Skill 未补填。")
    im.save(path,pnginfo=metadata)


def export_html(document,path,*,compact=False,od="after"):
    coverage=turn_coverage(document["rows"])
    gaps=coverage["gaps"]
    headers=HEADERS.copy()
    headers[7]="行动前 OD" if od=="before" else "行动后 OD"
    body=[]
    for row in display_rows(document,compact):
        body.append("<tr>"+"".join(f"<td>{escape(str(v))}</td>" for v in fields(row,od=od,coverage=coverage))+"</tr>")
    gap_banner=''
    if gaps:
        gap_banner=f'<aside class="coverage-gap" role="note"><strong>检测到 {coverage["missing_turn_count"]} 个缺失回合，以下区间未识别</strong><ul>'+''.join('<li>'+escape(gap_text(gap))+'</li>' for gap in gaps)+'</ul><p>这些区间未记录可见选招，角色与 Skill 未补填。请回看视频核对 AUTO、遮挡或剪辑。</p></aside>'
    html=f"""<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>{escape(document['title'])}</title><style>
body{{font:16px 'Microsoft YaHei',sans-serif;background:#f4f7fb;color:#28364e;padding:28px}} h1{{font-size:24px}}table{{border-collapse:collapse;background:white;width:100%}}td,th{{border:1px solid #d8dfea;padding:12px;text-align:center}}th{{background:#edf1f7}}td:nth-child(2),td:nth-child(3){{background:{COLORS[0]}}}td:nth-child(4),td:nth-child(5){{background:{COLORS[1]}}}td:nth-child(6),td:nth-child(7){{background:{COLORS[2]}}}td:nth-child(8){{color:#478dc9}}p{{color:#64748b}}.coverage-gap{{background:#fff0d3;border:1px solid #e8c98f;border-radius:8px;padding:16px 20px;margin:18px 0;color:#945b20}}.coverage-gap p{{color:inherit;margin-bottom:0}}.coverage-gap li{{margin:6px 0}}@media print{{body{{padding:0;background:white}}}}
</style><h1>{escape(document['title'])}</h1>{gap_banner}<p>每行对应一次实际行动确认。完整 Skill、原始 OCR 与 Screenshot 保存在配套文件中。</p><table><thead><tr>{''.join('<th>'+escape(h)+'</th>' for h in headers)}</tr></thead><tbody>{''.join(body)}</tbody></table><p>OD 为视频中的可见百分比。标记「待核对」的结果需要检查截图。</p></html>"""
    Path(path).write_text(html,encoding="utf-8")


def export_all(document,folder,*,compact=False,od="after"):
    folder=Path(folder)
    target=folder/"exports"/("compact" if compact else "detailed")
    target.mkdir(parents=True,exist_ok=True)
    for function,extension in [(export_csv,"csv"),(export_excel,"xlsx"),(export_png,"png"),(export_html,"html")]:
        function(document,target/f"行动轴.{extension}",compact=compact,od=od)
    (target/"行动轴.json").write_text(json.dumps({**document,"coverage":turn_coverage(document["rows"])},ensure_ascii=False,indent=2),encoding="utf-8")
    return target


def export_bundle(document,folder,*,compact=False,od="after"):
    folder=Path(folder)
    target=export_all(document,folder,compact=compact,od=od)
    destination=target/"行动轴与核对资料.zip"
    with zipfile.ZipFile(destination,"w",zipfile.ZIP_DEFLATED) as archive:
        for file in target.iterdir():
            if file.is_file() and file!=destination:archive.write(file,"exports/"+file.name)
        for relative in ("axis.json","axis.automatic.json","bootstrap.jpg","roster.jpg","capture.json"):
            if (folder/relative).is_file():archive.write(folder/relative,relative)
        for directory in ("frames","avatars","od"):
            for file in (folder/directory).glob("*.jpg"):archive.write(file,str(file.relative_to(folder)))
        for name in ("reference.png","validation.md"):
            if (folder/name).is_file():archive.write(folder/name,name)
    return destination
