"""Every exported format discloses missing Turns without inventing Actions."""
import copy
import csv
import json
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import load_workbook
from PIL import Image, ImageFont

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from axis.export import HEADERS, export_all, export_csv, export_excel, export_html, export_png, font_path


def has_cjk_font():
    try:
        ImageFont.truetype(font_path(), 18)
        return True
    except (OSError, ValueError):
        return False


def export_document():
    team=[dict(id=index,name=f'Character {index}',alias=f'角色 {index}') for index in range(1,7)]
    rows=[]
    for index,(turn,when) in enumerate([(11,10.),(22,20.)],1):
        actions=[dict(slot=slot,character_id=slot,character=f'角色 {slot}',character_full=f'Character {slot}',
                      skill=f'Skill {turn}.{slot}',skill_full=f'Full Skill {turn}.{slot}') for slot in range(1,4)]
        rows.append(dict(id=f'row-{index}',index=index,turn=turn,label=str(turn),kind='normal',time=when,
                         confirm_time=when+.25,phase_start=when-1,actions=actions,positions=[1,2,3,4,5,6],needs_review=False,
                         reviewed=False,od_before=10.,od_after=20.,notes=[]))
    return dict(title='Synthetic Gap Regression',team=team,rows=rows,video=dict(name='synthetic.mp4'),
                settings=dict(start=0),warnings=[])


class GapExportTests(unittest.TestCase):
    def test_csv_keeps_action_rows_and_columns_and_discloses_gap_in_notes(self):
        document=export_document();original=copy.deepcopy(document)
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'axis.csv'
            export_csv(document,path)
            with path.open(encoding='utf-8-sig',newline='') as stream:rows=list(csv.reader(stream))
        self.assertEqual(rows[0],HEADERS)
        self.assertEqual(len(rows),3)
        self.assertTrue(all(len(row)==10 for row in rows))
        self.assertEqual([row[0] for row in rows[1:]],['11','22'])
        self.assertEqual(rows[2][2],'Skill 22.1')
        self.assertIn('12–21',rows[2][-1])
        self.assertIn('未识别',rows[2][-1])
        self.assertIn('00:10',rows[2][-1]);self.assertIn('00:19',rows[2][-1])
        self.assertEqual(document,original)

    def test_html_recomputes_stale_coverage_and_escapes_known_actions(self):
        document=export_document()
        document['coverage']=dict(status='complete',missing_turn_count=0,gaps=[])
        document['rows'][0]['actions'][0]['skill']='<script>unsafe</script>'
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'axis.html';export_html(document,path)
            html=path.read_text(encoding='utf-8')
        self.assertIn('<aside class="coverage-gap"',html)
        self.assertLess(html.index('<aside'),html.index('<table>'))
        self.assertIn('12–21',html);self.assertIn('未识别',html)
        self.assertEqual(html.count('<tr>'),3)
        self.assertIn('&lt;script&gt;unsafe&lt;/script&gt;',html)
        self.assertNotIn('<script>',html)

    def test_excel_discloses_gap_without_changing_known_actions(self):
        document=export_document()
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'axis.xlsx';export_excel(document,path)
            workbook=load_workbook(path)
            try:
                self.assertIn('缺失回合',workbook.sheetnames)
                self.assertEqual(workbook['缺失回合']['A2'].value,'12–21')
                self.assertEqual(workbook['缺失回合']['E2'].value,'未识别')
                self.assertIn('缺失回合',workbook['行动轴']['A1'].value)
                self.assertIn('12–21',workbook['行动轴']['J4'].value)
                self.assertEqual(workbook['行动轴'].max_row,4)
                self.assertEqual(workbook['行动轴']['C4'].value,'Skill 22.1')
                self.assertEqual(workbook['完整 Skill 名称']['C4'].value,'Full Skill 22.1')
            finally:workbook.close()

    @unittest.skipUnless(has_cjk_font(),'PNG rendering requires an installed CJK Font')
    def test_png_contains_visible_gap_banner_and_text_metadata(self):
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'axis.png';export_png(export_document(),path)
            with Image.open(path) as image:
                self.assertIn('12–21',image.info['Coverage'])
                self.assertIn('未识别',image.info['Coverage'])
                self.assertEqual(image.getpixel((5,80)),(255,240,211))
                self.assertGreater(image.height,125+2*53+92)

    @unittest.skipUnless(has_cjk_font(),'Complete export includes PNG rendering with a CJK Font')
    def test_json_computes_legacy_coverage_without_fabricating_actions(self):
        document=export_document();original=copy.deepcopy(document)
        with tempfile.TemporaryDirectory() as temporary:
            target=export_all(document,temporary,compact=True)
            saved=json.loads((target/'行动轴.json').read_text(encoding='utf-8'))
        self.assertEqual(saved['coverage']['missing_turn_count'],10)
        self.assertEqual(saved['coverage']['gaps'][0]['missing_turns'],list(range(12,22)))
        self.assertEqual(saved['rows'],original['rows'])
        self.assertEqual(document,original)

    def test_corrected_rows_remove_stale_gap_metadata_from_exports(self):
        document=export_document()
        document['rows'][1].update(turn=12,label='12')
        document['coverage']=dict(status='gaps_detected',missing_turn_count=10,gaps=[dict(missing_turns=list(range(12,22)))])
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'axis.html';export_html(document,path)
            self.assertNotIn('<aside',path.read_text(encoding='utf-8'))
            path=Path(temporary)/'axis.xlsx';export_excel(document,path)
            workbook=load_workbook(path)
            try:self.assertNotIn('缺失回合',workbook.sheetnames)
            finally:workbook.close()

    def test_legacy_missing_phase_time_is_disclosed_without_a_made_up_timestamp(self):
        document=export_document();document['rows'][1].pop('phase_start')
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'axis.xlsx';export_excel(document,path)
            workbook=load_workbook(path)
            try:
                self.assertIn('时间未记录',workbook['缺失回合']['D2'].value)
                self.assertIn('00:10',workbook['缺失回合']['D2'].value)
            finally:workbook.close()


if __name__=='__main__':
    unittest.main()
