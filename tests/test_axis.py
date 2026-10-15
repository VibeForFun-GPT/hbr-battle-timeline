"""验证 Turn 规则、Character Identity、OCR 和本地 Export。"""
import copy
import csv
from http.client import HTTPConnection
from http.server import ThreadingHTTPServer
import json
from pathlib import Path
import sys
import tempfile
import threading
import unittest
import zipfile

import cv2
import numpy as np
from openpyxl import load_workbook
from PIL import Image
from PIL import ImageDraw,ImageFont
from unittest.mock import patch

HERE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(HERE))
from app import Store, make_handler, parse_time, validate_document, validate_profile
from axis.export import export_bundle, font_path
from axis.model import annotate_rows, build_team, compact_rows, format_skill_sequence, resolve_character, resolve_skill
from axis.pipeline import analyze, is_od_menu_transition, save_json
from axis.vision import Avatars, DEFAULT_LAYOUT, Reader, Video, parse_counter, read_image
from _support import cjk_font_ready, private_frames_ready, sample_profile_ready


def action(ident,skill='A'):
    return dict(character_id=ident,character=f'角色 {ident}',character_full=f'Name {ident}',slot=ident,skill=skill,skill_full='Normal Attack' if skill=='A' else skill,skill_raw=skill,confidence=.9)


def row(turn,bonus=None,ids=(1,2,3),skill='A',**extra):
    return dict(turn=turn,bonus=bonus,time=1.,confirm_time=1.2,actions=[action(i,skill) for i in ids],needs_review=False,reviewed=True,od_before=0.,od_after=5.,**extra)


def document():
    rows=annotate_rows([row(1),row(1,ids=(3,1,4),skill='Buff')])
    return dict(schema_version=1,title='Regression',created_at='2026-10-04T00:00:00+08:00',team=[dict(id=i,name=f'Name {i}',alias=f'角色 {i}',avatar=f'avatars/{i}.jpg') for i in range(1,7)],settings=dict(start=0,layout=copy.deepcopy(DEFAULT_LAYOUT),aliases={}),rows=rows,video=dict(name='tiny.mp4',path='',duration=10,width=1920,height=1080),warnings=[])


class ModelTests(unittest.TestCase):
    def test_pre_post_and_bonus(self):
        rows=annotate_rows([row(1),row(2,[1,2],od_level=2),row(2,[1,2]),row(2,[2,2]),row(2),row(2,[1,3],od_level=5,od_activation=True)])
        self.assertEqual([r['kind'] for r in rows],['normal','pre_od','extra','od','normal','post_od'])
        self.assertEqual(rows[-1]['label'],'后置OD5')
        self.assertEqual(rows[-1]['bonus_text'],'BONUS 1/3')

    def test_repeated_normal_turn_is_extra(self):
        rows=annotate_rows([row(1),row(1),row(2)])
        self.assertEqual([r['label'] for r in rows],['1','追加','2'])

    def test_compaction_tracks_identity_and_review(self):
        rows=annotate_rows([row(1),row(1,ids=(3,1,4),skill='Buff')])
        rows[1].update(reviewed=False,needs_review=True,od_after=42.)
        compact=compact_rows(rows)
        self.assertEqual([a['skill'] for a in compact[0]['actions']],['A + Buff','A','A + Buff'])
        self.assertEqual(compact[0]['extra_actions'][0]['character_id'],4)
        self.assertEqual(compact[0]['od_after'],42.)
        self.assertFalse(compact[0]['reviewed'])
        self.assertEqual(rows[0]['actions'][0]['skill'],'A')
        self.assertEqual(format_skill_sequence(['代码','代码','A']),'代码×2 + A')

    def test_counter_ocr_repair(self):
        self.assertEqual(parse_counter('TURN7130','TURN'),[7,30])
        self.assertEqual(parse_counter('BONUS 1 / 2','BONUS'),[1,2])
        self.assertIsNone(parse_counter('TURN 32/30','TURN'))
        self.assertIsNone(parse_counter('damage 1/30','TURN'))

    def test_unknown_skill_and_full_name(self):
        full,short,confidence,_=resolve_skill('重點保養',{'重點保養':'sp'})
        self.assertEqual((full,short),('重點保養','sp'))
        self.assertEqual(resolve_skill('未知新增技能')[0:2],('未知新增技能','未知新增技能'))

    def test_similar_new_skills_are_not_replaced(self):
        for text in ['血腥之舞蹈','重點保養','星火之歌+','追加回合之歌','Healing Wave']:
            full,short,confidence,_=resolve_skill(text,{'血腥之舞':'血腥','星火之歌':'火歌'})
            self.assertEqual((full,short),(text,text))
            self.assertEqual(confidence,1.)

    def test_names_are_placeholders_or_arbitrary_input(self):
        team=build_team()
        self.assertEqual([p['alias'] for p in team],[f'角色 {i}' for i in range(1,7)])
        self.assertTrue(all(p['name_is_placeholder'] for p in team))
        names=['队员甲','任意人乙','New Character','自定义昵称','角色丙','其它队员']
        named=build_team([{'alias':name} for name in names])
        self.assertEqual([p['name'] for p in named],names)
        self.assertEqual(resolve_character('白河另一人')['name'],'白河另一人')
        self.assertEqual(validate_profile({})['name_mode'],'auto')
        self.assertEqual(validate_profile({'name_mode':'placeholder'})['name_mode'],'placeholder')
        self.assertEqual(validate_profile({'team':[{'name':n} for n in names]})['name_mode'],'profile')

    def test_manual_name_updates_every_action_for_identity(self):
        d=document();old=copy.deepcopy(d);d['team'][0].update(alias='任意昵称',name='完全不同的人')
        changed=validate_document(d,old)
        targets=[a for r in changed['rows'] for a in r['actions'] if a['character_id']==1]
        self.assertEqual(len(targets),2)
        self.assertTrue(all(a['character']=='任意昵称' and a['character_full']=='完全不同的人' for a in targets))

    def test_triple_attack_before_post_od_is_preserved(self):
        normal=row(9);normal.update(phase_start=0,time=10,confirm_time=10.25)
        following=row(9,[1,3]);following.update(time=20)
        self.assertFalse(is_od_menu_transition(normal,following,[{'time':15}]))
        transition=row(9);transition.update(phase_start=443.633,time=443.9,confirm_time=444.167)
        self.assertTrue(is_od_menu_transition(transition,following,[{'time':444.433}]))

    def test_validation(self):
        self.assertEqual(parse_time('01:13'),73.)
        with self.assertRaises(ValueError):parse_time('NaN')
        with self.assertRaises(ValueError):validate_profile({'layout':{'crop':[0,0,0,1]}})
        with self.assertRaises(ValueError):validate_profile({'layout':{'portrait_centers':[[0,0]]*5}})
        d=document();original=copy.deepcopy(d);d['rows'][0]['actions'][1]['character_id']=1
        with self.assertRaises(ValueError):validate_document(d,original)
        d=document();d['video']['path']='injected'
        self.assertEqual(validate_document(d,original)['video']['path'],'')


class ExportTests(unittest.TestCase):
    @unittest.skipUnless(cjk_font_ready(), 'PNG integration requires an installed CJK Font (Microsoft YaHei or Noto Sans CJK)')
    def test_all_exports_keep_full_names_and_safe_text(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder=Path(temporary);d=document();d['rows'][0]['actions'][0]['skill']='=1+1'
            (folder/'axis.json').write_text(json.dumps(d),encoding='utf8')
            bundle=export_bundle(d,folder,compact=True)
            target=bundle.parent
            workbook=load_workbook(target/'行动轴.xlsx')
            self.assertIn('完整 Skill 名称',workbook.sheetnames)
            self.assertIn('识别依据',workbook.sheetnames)
            self.assertEqual(workbook['行动轴']['C3'].data_type,'s')
            with (target/'行动轴.csv').open(encoding='utf-8-sig',newline='') as stream:
                values=list(csv.reader(stream))
            self.assertTrue(values[1][2].startswith("'="))
            with Image.open(target/'行动轴.png') as image:self.assertGreater(image.width,1000)
            self.assertEqual(json.loads((target/'行动轴.json').read_text(encoding='utf8'))['rows'][0]['actions'][0]['skill_full'],'Normal Attack')
            with zipfile.ZipFile(bundle) as archive:self.assertIn('axis.json',archive.namelist())


class ServerTests(unittest.TestCase):
    def test_atomic_save_while_a_reader_holds_the_file(self):
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'job.json'
            save_json(path,{'fraction':0})
            reader=path.open('rb')
            self.assertTrue(reader.read())
            release=threading.Timer(.12,reader.close)
            release.start()
            try:
                save_json(path,{'fraction':1,'status':'done'})
            finally:
                release.join();reader.close()
            self.assertEqual(json.loads(path.read_text(encoding='utf8')),{'fraction':1,'status':'done'})
            self.assertFalse(path.with_suffix('.json.tmp').exists())

    def test_local_api_and_video_range(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder=Path(temporary);sample=folder/'sample-final';sample.mkdir()
            video=sample/'tiny.mp4';video.write_bytes(b'0123456789')
            d=document();d['video']['path']=str(video)
            (sample/'axis.json').write_text(json.dumps(d),encoding='utf8')
            server=ThreadingHTTPServer(('127.0.0.1',0),make_handler(Store(folder),'test'))
            worker=threading.Thread(target=server.serve_forever,daemon=True);worker.start()
            conn=HTTPConnection('127.0.0.1',server.server_port,timeout=10)
            try:
                conn.request('GET','/test/video/sample',headers={'Range':'bytes=2-5'})
                response=conn.getresponse();self.assertEqual(response.status,206);self.assertEqual(response.read(),b'2345')
                conn.request('GET','/test/video/sample',headers={'Range':'bytes=-3'})
                response=conn.getresponse();self.assertEqual(response.read(),b'789')
                conn.request('GET','/test/video/sample',headers={'Range':'bytes=99-'})
                response=conn.getresponse();self.assertEqual(response.status,416);response.read()
                conn.request('GET','/test/file/sample/../../outside.json')
                response=conn.getresponse();self.assertEqual(response.status,400);response.read()
                conn.request('POST','/test/api/save/sample',json.dumps({'document':d}),{'Content-Type':'application/json','Origin':'https://example.invalid'})
                response=conn.getresponse();self.assertEqual(response.status,400);response.read()
                conn.request('GET','/wrong/api/jobs')
                response=conn.getresponse();self.assertEqual(response.status,400);response.read()
            finally:conn.close();server.shutdown();server.server_close();worker.join()


class VideoIntegrationTests(unittest.TestCase):
    def test_synthetic_video_decode_crop_and_seeking(self):
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'synthetic.avi'
            writer=cv2.VideoWriter(str(path),cv2.VideoWriter_fourcc(*'MJPG'),10,(320,180))
            self.assertTrue(writer.isOpened())
            try:
                for index in range(20):
                    frame=np.full((180,320,3),(0,0,220) if index<10 else (220,0,0),dtype=np.uint8)
                    frame[:,:70]=255
                    frame[:,250:]=255
                    writer.write(frame)
            finally:writer.release()
            video=Video(path,{'crop':[.25,.2,.75,.8]})
            try:
                self.assertAlmostEqual(video.duration,2.,places=1)
                first=video.frame(0)
                later=video.frame(1.2)
                self.assertEqual(first.shape,(1080,1920,3))
                self.assertEqual(later.shape,first.shape)
                self.assertGreater(float(first[:,:,2].mean()),210)
                self.assertLess(float(first[:,:,0].mean()),10)
                self.assertGreater(float(later[:,:,0].mean()),210)
                self.assertLess(float(later[:,:,2].mean()),10)
                samples=list(video.sample(.8,1.5,rate=4))
                self.assertGreaterEqual(len(samples),3)
                self.assertTrue(all(.8<=when<1.5 and frame.shape==first.shape for when,frame in samples))
            finally:video.close()

    @unittest.skipUnless(private_frames_ready() and sample_profile_ready() and cjk_font_ready(), 'Optional real-media integration requires all 7 local Frames, sample Profile and a CJK Font')
    def test_real_frames_at_720p(self):
        fixtures=HERE/'tests'/'fixtures'
        profile=json.loads((HERE/'examples'/'sample-profile.json').read_text(encoding='utf8'))
        with tempfile.TemporaryDirectory() as temporary:
            folder=Path(temporary);path=folder/'compressed.mp4'
            writer=cv2.VideoWriter(str(path),cv2.VideoWriter_fourcc(*'mp4v'),10,(1280,720))
            self.assertTrue(writer.isOpened())
            scenes=[('bootstrap.jpg',8)]
            for n in range(1,4):scenes += [(f'select-{n}.jpg',20),(f'confirm-{n}.jpg',6),(None,18)]
            for filename,count in scenes:
                frame=read_image(fixtures/filename) if filename else np.zeros((720,1280,3),np.uint8)
                for _ in range(count):writer.write(frame)
            writer.release()
            d=analyze(path,folder/'output',start=0,sample_rate=4,profile=profile)
            self.assertEqual([r['label'] for r in d['rows']],['1','2','追加'])
            self.assertEqual([a['character'] for a in d['rows'][0]['actions']],['可怜','三乡','diva'])
            self.assertEqual([a['skill'] for a in d['rows'][0]['actions']],['sp','甜点','sp'])
            self.assertEqual([a['character'] for a in d['rows'][1]['actions']],['可怜','白河','yuki'])
            self.assertEqual(d['rows'][1]['actions'][0]['skill'],'血腥+')
            self.assertTrue((folder/'output'/'axis.automatic.json').exists())
            video=Video(path);self.assertEqual(video.frame(0).shape,(1080,1920,3));video.close()

    @unittest.skipUnless(private_frames_ready() and sample_profile_ready() and cjk_font_ready(), 'Optional real-media integration requires all 7 local Frames, sample Profile and a CJK Font')
    def test_unseen_avatars_skills_and_automatic_name_fallback(self):
        fixtures=HERE/'tests'/'fixtures'
        generator=np.random.default_rng(715)
        faces=[cv2.resize(generator.integers(20,240,(18,19,3),dtype=np.uint8),(78,72),interpolation=cv2.INTER_NEAREST) for _ in range(6)]
        skills=['冰之波纹','星火之歌','勇气之证']
        orders={'bootstrap.jpg':[0,1,2,3,4,5],'select-1.jpg':[0,3,2,1,4,5],'select-2.jpg':[0,4,1,2,3,5],'select-3.jpg':[0,4,1,2,3,5]}
        font=ImageFont.truetype(font_path(),34)
        with tempfile.TemporaryDirectory() as temporary:
            folder=Path(temporary);path=folder/'new-team.mp4'
            writer=cv2.VideoWriter(str(path),cv2.VideoWriter_fourcc(*'mp4v'),10,(1280,720))
            self.assertTrue(writer.isOpened())
            scenes=[(None,60),('bootstrap.jpg',8)]
            for n in range(1,4):scenes += [(f'select-{n}.jpg',20),(f'confirm-{n}.jpg',6),(None,18)]
            for filename,count in scenes:
                if filename:
                    frame=cv2.resize(read_image(fixtures/filename),(1920,1080))
                    order=orders.get(filename,orders[filename.replace('confirm','select')]) if filename!='bootstrap.jpg' else orders[filename]
                    for slot,(x,y) in enumerate(DEFAULT_LAYOUT['portrait_centers']):frame[y-15:y+57,x-39:x+39]=faces[order[slot]]
                    image=Image.fromarray(cv2.cvtColor(frame,cv2.COLOR_BGR2RGB));draw=ImageDraw.Draw(image)
                    for text,box in zip(skills,DEFAULT_LAYOUT['skill_boxes']):
                        draw.rectangle(box,fill='#27242d');draw.text(((box[0]+box[2])/2,(box[1]+box[3])/2),text,font=font,fill='white',anchor='mm')
                    frame=cv2.resize(cv2.cvtColor(np.asarray(image),cv2.COLOR_RGB2BGR),(1280,720))
                else:frame=np.zeros((720,1280,3),np.uint8)
                for _ in range(count):writer.write(frame)
            writer.release()
            with patch.object(Reader,'roster',side_effect=AssertionError('Placeholder Mode 不应读取角色配置')):
                d=analyze(path,folder/'output',sample_rate=4)
            self.assertEqual(d['settings']['name_mode'],'auto')
            self.assertEqual([p['alias'] for p in d['team']],[f'角色 {i}' for i in range(1,7)])
            self.assertEqual([r['label'] for r in d['rows']],['1','2','追加'])
            self.assertEqual([a['character_id'] for a in d['rows'][0]['actions']],[1,4,3])
            self.assertEqual([a['character_id'] for a in d['rows'][1]['actions']],[1,5,2])
            self.assertEqual([a['skill'] for a in d['rows'][0]['actions']],skills)
            self.assertTrue(all(a['skill_confidence']>.85 for a in d['rows'][0]['actions']))
            self.assertFalse(any('名称需要核对' in n for r in d['rows'] for n in r['notes']))

    def test_direct_ocr_fallback_retains_a_new_skill(self):
        reader=object.__new__(Reader);reader.layout=DEFAULT_LAYOUT;reader.aliases={}
        reader.ocr=lambda *a,**k:(None,None)
        reader.lines=lambda patches:[('星火之歌',.98),('',.3),('勇气之证…',.92)]
        actions=reader.skills(np.zeros((1080,1920,3),np.uint8))
        self.assertEqual(actions[0]['skill'],'星火之歌')
        self.assertEqual(actions[1]['skill'],'A')
        self.assertEqual(actions[2]['skill'],'勇气之证…')
        self.assertTrue(actions[2]['skill_truncated'])


if __name__=='__main__':unittest.main(verbosity=2)
